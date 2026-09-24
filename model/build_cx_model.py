"""Rebuild the Revenue Build and DCF tabs of CX_VIA_Model.xlsx around Via's theses.

Keeps the workbook's formatting (styles are copied from the template's own rows) and the
Capital IQ supporting tabs. Run:  python3 model/build_cx_model.py <in.xlsx> <out.xlsx>
"""
import sys
from copy import copy
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font

SRC, OUT = sys.argv[1], sys.argv[2]
BBG = sys.argv[3] if len(sys.argv) > 3 else "model/Consensus_Projections_BBG_quarterly.xlsx"
wb = openpyxl.load_workbook(SRC)

# ---------------------------------------------------------------- Bloomberg consensus tab
# Quarterly Bloomberg consensus (Q3'25A-Q4'27E) copied in as values, plus FY26E/FY27E totals.
src = openpyxl.load_workbook(BBG, data_only=True).worksheets[0]
if "Consensus (BBG)" in wb.sheetnames:
    del wb["Consensus (BBG)"]
cb = wb.create_sheet("Consensus (BBG)", index=wb.sheetnames.index("Supporting Tabs >>>") + 1)
for row in src.iter_rows():
    for cell in row:
        if cell.value is not None:
            n = cb.cell(row=cell.row, column=cell.column, value=cell.value)
            n.font = Font(name="Arial", size=9, bold=cell.font.b)
            if isinstance(cell.value, (int, float)):
                n.number_format = "#,##0.00"
cb.column_dimensions["A"].width = 34
cb["P3"], cb["Q3"] = "FY2026E", "FY2027E"
cb["P4"], cb["Q4"] = "Sum of Q1-Q4", "Sum of Q1-Q4"
for rr in (7, 10, 22):                       # revenue, adj. EBITDA, gross profit
    cb[f"P{rr}"] = f"=SUM(G{rr}:J{rr})"
    cb[f"Q{rr}"] = f"=SUM(K{rr}:N{rr})"
cb["P23"], cb["Q23"] = "=P22/P7", "=Q22/Q7"   # gross margin
for ref in ["P3", "Q3", "P4", "Q4"]:
    cb[ref].font = Font(name="Arial", size=9, bold=True)
for rr in (7, 10, 22, 23):
    for c in "PQ":
        cb[f"{c}{rr}"].font = Font(name="Arial", size=9, bold=True)
        cb[f"{c}{rr}"].number_format = "0.0%" if rr == 23 else "#,##0.0"
cb["P1"] = "Source: Bloomberg consensus (VIA US Equity), quarterly. FY columns sum the four quarters (FY26 includes Q1-Q2 2026 actuals)."
cb["P1"].font = Font(name="Arial", size=9, italic=True)
BBG_COL = {"F": "P", "G": "Q"}               # FY26E, FY27E; FY28E falls back to CapIQ
rb, dc, wacc = wb["Revenue Build"], wb["DCF"], wb["WACC Calculations"]

HIST, FCST = ["C", "D", "E"], ["F", "G", "H", "I", "J"]
ALL = HIST + FCST
IS_COL = {"C": "B", "D": "C", "E": "D"}          # Annual IS column for each historical year
KS_COL = {"F": "F", "G": "G", "H": "H"}          # Key Stats consensus columns (FY26-28E)
USD, USD1, PCT, NUM1, MULT = "#,##0;\\(#,##0\\);\\-", "#,##0.0;\\(#,##0.0\\);\\-", "0.0%", "#,##0.0", "0.0x"
BLUE, GREEN, BLACK = "FF0000FF", "FF008000", "FF000000"


def note(cell, text):
    cell.comment = Comment(text, "Model")


# ---------------------------------------------------------------- style prototypes
def grab(ws, r):
    return {c: copy(ws[f"{c}{r}"]._style) for c in ["B"] + ALL}

P = {k: grab(rb, r) for k, r in dict(sub=4, line=5, ital=7, tot=8, gr=9, cons=10, consgr=11,
                                      band=13, subhdr=14, drv=15, drvgr=16).items()}


def style(r, proto, fmt=None, color=None, cols=ALL, italic=None, bold=None):
    for c in ["B"] + cols:
        cell = rb[f"{c}{r}"]
        cell._style = copy(P[proto][c])
        if c != "B":
            if fmt: cell.number_format = fmt
            if color:
                f = copy(cell.font); f.color = color; cell.font = f
        if italic is not None or bold is not None:
            f = copy(cell.font)
            if italic is not None: f.i = italic
            if bold is not None: f.b = bold
            cell.font = f


def setfont(ws, ref, color):
    f = copy(ws[ref].font); f.color = color; ws[ref].font = f


# ---------------------------------------------------------------- clear Revenue Build body
for mr in list(rb.merged_cells.ranges):
    if mr.min_row >= 4:
        rb.unmerge_cells(str(mr))
for row in rb.iter_rows(min_row=4, max_row=70, min_col=2, max_col=10):
    for cell in row:
        cell.value = None
        cell._style = copy(rb["L4"]._style)
        cell.comment = None
rb.column_dimensions["B"].width = 44

R = {}  # row map


def put_row(key, r, label):
    R[key] = r
    rb[f"B{r}"] = label


# ---------------------------------------------------------------- revenue block
put_row("sw", 4, "Software (platform)");                 style(4, "line", USD)
put_row("svc", 5, "Services");                           style(5, "sub", USD, bold=True)
put_row("mp", 6, "  Microtransit & paratransit");        style(6, "line", USD)
put_row("nw", 7, "  Network deals");                     style(7, "line", USD)
put_row("ot", 8, "One-time (implementation, consulting)"); style(8, "ital", USD)
put_row("tot", 9, "Total net revenue");                  style(9, "tot", USD)
put_row("tg", 10, "   % growth");                        style(10, "gr", PCT, italic=True)
put_row("cons", 11, "Consensus (BBG; FY28 CapIQ)");                style(11, "cons", USD)
put_row("cg", 12, "   % growth");                        style(12, "consgr", PCT, italic=True)
put_row("vs", 13, "   Model vs consensus");              style(13, "consgr", PCT, italic=True)
put_row("ss", 14, "Services % of revenue");              style(14, "ital", PCT, italic=True)
put_row("nws", 15, "Network deals % of revenue");        style(15, "ital", PCT, italic=True)
put_row("nwg", 16, "Network deals share of revenue growth"); style(16, "ital", PCT, italic=True)

# ---------------------------------------------------------------- drivers (one per stream)
put_row("dh", 18, "DRIVERS");                            style(18, "band", cols=ALL)
for c in ALL: rb[f"{c}18"].value = None
put_row("sw_g", 19, "Software revenue % growth");        style(19, "drvgr", PCT)
put_row("mp_g", 20, "Microtransit & paratransit % growth (net of downsell)"); style(20, "drvgr", PCT)
put_row("new", 21, "Network: new contract value launched ($mm)"); style(21, "drvgr", USD1)
put_row("rr", 22, "Network: contract run-rate, year-end ($mm)"); style(22, "drv", USD1)
put_row("otp", 23, "One-time revenue % of total");       style(23, "drvgr", PCT)

put_row("h_gp", 25, "GROSS MARGIN BY STREAM");           style(25, "subhdr")
put_row("gm_sw", 26, "Software");                        style(26, "drvgr", PCT)
put_row("gm_mp", 27, "Microtransit & paratransit");      style(27, "drvgr", PCT)
put_row("gm_nw", 28, "Network deals");                   style(28, "drvgr", PCT)
put_row("gm_ot", 29, "One-time");                        style(29, "drvgr", PCT)
put_row("gp", 30, "Gross profit ($mm)");                 style(30, "tot", USD)
put_row("gm", 31, "   Gross margin");                    style(31, "gr", PCT, italic=True)
put_row("gmc", 32, "   Consensus gross margin (BBG; FY28 CapIQ)"); style(32, "consgr", PCT, italic=True)

# historical cells in input rows are black formulas (or blue hardcodes set below)
for key in ["sw_g", "mp_g", "new", "otp", "gm_sw", "gm_mp", "gm_nw", "gm_ot"]:
    for c in HIST:
        rb[f"{c}{R[key]}"]._style = copy(P["drv"][c]); rb[f"{c}{R[key]}"].number_format = rb[f"F{R[key]}"].number_format
        setfont(rb, f"{c}{R[key]}", BLACK)

r = R
# ---- historical inputs (blue) ----
for c, v in zip(HIST, [57.9, 78.6, 101.1]):
    rb[f"{c}{r['sw']}"] = v; setfont(rb, f"{c}{r['sw']}", BLUE)
note(rb[f"E{r['sw']}"], "ESTIMATE. Via reports all recurring revenue as one 'subscription' line. Software ~24% of recurring revenue, services ~76% (margin back-out in VIA_software_services_mix.xlsx; Bleecker 72% as published).")
for c, v in zip(HIST, [0.0, 11.1, 23.5]):
    rb[f"{c}{r['rr']}"] = v; setfont(rb, f"{c}{r['rr']}", BLUE)
note(rb[f"D{r['rr']}"], "Sioux Falls: 2024 NTD purchased transportation paid $11.07M (Via's first network deal).")
note(rb[f"E{r['rr']}"], "Sioux Falls ~$11.4M + Mobile $12.1M/yr (from Oct 2025).")
for c, v in zip(HIST, [0.0, 11.1, 14.4]):
    rb[f"{c}{r['nw']}"] = v; setfont(rb, f"{c}{r['nw']}", BLUE)
note(rb[f"E{r['nw']}"], "ESTIMATE. Sioux Falls ~$11.4M + one quarter of Mobile ($12.1M/4).")
for c in HIST:
    rb[f"{c}{r['otp']}"] = 0.03; setfont(rb, f"{c}{r['otp']}", BLUE)
note(rb[f"E{r['otp']}"], "10-K: one-time revenue 3% of total in FY24 and FY25 (97% recurring). FY23 assumed the same.")

# ---- forecast inputs (blue) ----
FIN = {
    "sw_g":  ([0.03, 0.11, 0.10, 0.09, 0.08], "Organic customer growth ~9% ex-Downtowner (thesis 3), ~+2%/yr price. FY26 +3%: 94 small Downtowner customers add little. Software fees face near-zero bids (Spare bid $0.01 at LA Metro)."),
    "mp_g":  ([0.28, 0.10, 0.08, 0.07, 0.06], "Pilots are budget-bound (thesis 3): expansion at existing customers less downsell (Jersey City cut ~half; federal stopgap / ARPA cliff), price +2%/yr as rebids reset rates. Growth shifts to network deals. FY26 +28% ties to 1H26 actuals ($263M revenue) plus the H2 ramp."),
    "new":   ([20.0, 40.0, 60.0, 70.0, 80.0], "The growth engine (thesis 1). FY26: Twin Cities MI (Apr), Rochester $14.6M (Sep), part of the 2026 wins. FY27: rest of the signed ~$70M book (~$27M) plus new wins. FY28+: the >$700M pipeline (doubled YoY) converting at ~8-11% a year, ~5-7 mid-sized cities at $10-15M each. Earned at half rate in the launch year."),
    "otp":   ([0.03] * 5, "3% (10-K FY24-25; 1H26 also 3%)."),
    "gm_sw": ([0.75] * 5, "Bleecker assumption; reported cost split implies ~73%."),
    "gm_mp": ([0.293] * 5, "Bleecker TaaS gross margin, arithmetic corrected and updated to 2026 contract pricing."),
    "gm_nw": ([0.25] * 5, "ASSUMPTION (undisclosed): operator economics, below the ~29% on microtransit. Labor is 72% of mid-sized bus system cost (NTD 2024); CFO: 'some accretive, some less.'"),
    "gm_ot": ([0.75] * 5, "Implementation / consulting."),
}
for key, (vals, txt) in FIN.items():
    for c, v in zip(FCST, vals):
        rb[f"{c}{r[key]}"] = v
    note(rb[f"F{r[key]}"], txt)

# ---- formulas ----
for i, c in enumerate(ALL):
    p = ALL[i - 1] if i else None
    if c in HIST:
        rb[f"{c}{r['tot']}"] = f"='Annual IS'!{IS_COL[c]}7"; setfont(rb, f"{c}{r['tot']}", GREEN)
        rb[f"{c}{r['ot']}"] = f"={c}{r['tot']}*{c}{r['otp']}"
        rb[f"{c}{r['mp']}"] = f"={c}{r['tot']}-{c}{r['ot']}-{c}{r['sw']}-{c}{r['nw']}"
        rb[f"{c}{r['cons']}"] = f"={c}{r['tot']}"
        rb[f"{c}{r['gp']}"] = f"='Annual IS'!{IS_COL[c]}10"; setfont(rb, f"{c}{r['gp']}", GREEN)
        rb[f"{c}{r['gmc']}"] = f"={c}{r['gm']}"
        if p:
            rb[f"{c}{r['sw_g']}"] = f"={c}{r['sw']}/{p}{r['sw']}-1"
            rb[f"{c}{r['mp_g']}"] = f"={c}{r['mp']}/{p}{r['mp']}-1"
            rb[f"{c}{r['new']}"] = f"={c}{r['rr']}-{p}{r['rr']}"
            rb[f"{c}{r['cg']}"] = f"={c}{r['cons']}/{p}{r['cons']}-1"
    else:
        rb[f"{c}{r['sw']}"] = f"={p}{r['sw']}*(1+{c}{r['sw_g']})"
        rb[f"{c}{r['mp']}"] = f"={p}{r['mp']}*(1+{c}{r['mp_g']})"
        rb[f"{c}{r['rr']}"] = f"={p}{r['rr']}+{c}{r['new']}"
        rb[f"{c}{r['nw']}"] = f"={p}{r['rr']}+0.5*{c}{r['new']}"
        rb[f"{c}{r['ot']}"] = f"=({c}{r['sw']}+{c}{r['svc']})*{c}{r['otp']}/(1-{c}{r['otp']})"
        rb[f"{c}{r['tot']}"] = f"={c}{r['sw']}+{c}{r['svc']}+{c}{r['ot']}"
        rb[f"{c}{r['gp']}"] = (f"={c}{r['sw']}*{c}{r['gm_sw']}+{c}{r['mp']}*{c}{r['gm_mp']}"
                               f"+{c}{r['nw']}*{c}{r['gm_nw']}+{c}{r['ot']}*{c}{r['gm_ot']}")
        if c in KS_COL:
            if c in BBG_COL:
                rb[f"{c}{r['cons']}"] = f"='Consensus (BBG)'!{BBG_COL[c]}7"
                rb[f"{c}{r['gmc']}"] = f"='Consensus (BBG)'!{BBG_COL[c]}23"
            else:
                rb[f"{c}{r['cons']}"] = f"='Key Stats'!{KS_COL[c]}16"
                rb[f"{c}{r['gmc']}"] = f"='Key Stats'!{KS_COL[c]}20"
            setfont(rb, f"{c}{r['cons']}", GREEN); setfont(rb, f"{c}{r['gmc']}", GREEN)
            rb[f"{c}{r['vs']}"] = f"={c}{r['tot']}/{c}{r['cons']}-1"
            rb[f"{c}{r['cg']}"] = f"={c}{r['cons']}/{p}{r['cons']}-1"
    rb[f"{c}{r['svc']}"] = f"={c}{r['mp']}+{c}{r['nw']}"
    if p: rb[f"{c}{r['tg']}"] = f"={c}{r['tot']}/{p}{r['tot']}-1"
    rb[f"{c}{r['ss']}"] = f"={c}{r['svc']}/{c}{r['tot']}"
    rb[f"{c}{r['nws']}"] = f"={c}{r['nw']}/{c}{r['tot']}"
    if p: rb[f"{c}{r['nwg']}"] = f"=({c}{r['nw']}-{p}{r['nw']})/({c}{r['tot']}-{p}{r['tot']})"
    rb[f"{c}{r['gm']}"] = f"={c}{r['gp']}/{c}{r['tot']}"

note(rb[f"C{r['tot']}"], "Historical total revenue links to Annual IS (CapIQ).")
note(rb[f"C{r['mp']}"], "Historical: residual = total - one-time - software - network.")
note(rb[f"F{r['nw']}"], "Opening run-rate + half of the contract value launched in the year.")

notes = [
    "Colour key: blue = hardcoded input, black = formula, green = link to another tab. Hover over cells with a red corner for sources.",
    "Thesis 1: network deals are bus-operator contracts priced at ~95-107% of the city's transit budget, so they carry a 25% gross margin vs ~29% on microtransit and 75% on software.",
    "Thesis 1: network deals become the growth engine (~a third of FY27 growth, ~25% of revenue by FY30), dragging gross margin toward operator levels. Thesis 3: microtransit/paratransit growth is net of budget-driven downsell and rebid price resets.",
    "Consensus: Bloomberg FY26-27E (Consensus (BBG) tab, sum of quarters); FY28E from CapIQ (Key Stats tab) because Bloomberg stops at Q4 2027. FY29-30 have no consensus.",
]
for k, t in enumerate(notes):
    rr = 34 + k
    rb.merge_cells(f"B{rr}:J{rr}")
    rb[f"B{rr}"] = t
    rb[f"B{rr}"].font = Font(name="Arial", size=9, italic=True)

# ================================================================= DCF
RBT = lambda c, key: f"'Revenue Build'!{c}{R[key]}"
labels = {9: "Hour-driven G&A (insurance, support)", 10: "Fixed opex (other G&A, S&M, R&D)",
          11: "Stock-based comp (in opex)", 32: "COGS as a % of Revenue",
          33: "Hour-driven G&A (% of services revenue)", 34: "Fixed opex growth",
          35: "SBC as a % of Revenue", 38: "Add back SBC to UFCF? (1 = yes, 0 = no)"}
for rr, t in labels.items():
    dc[f"B{rr}"] = t

for i, c in enumerate(ALL):
    p = ALL[i - 1] if i else None
    hist = c in HIST
    ic = IS_COL.get(c)
    dc[f"{c}4"] = f"={RBT(c, 'tot')}"; setfont(dc, f"{c}4", GREEN)
    if hist:
        dc[f"{c}6"] = f"='Annual IS'!{ic}9"
        dc[f"{c}9"] = f"={RBT(c, 'svc')}*{c}33"
        dc[f"{c}11"] = f"='Annual IS'!{ic}93-'Annual IS'!{ic}89"
        dc[f"{c}10"] = f"='Annual IS'!{ic}17-{c}9-{c}11"
        dc[f"{c}15"] = f"='Annual IS'!{ic}33"
        dc[f"{c}18"] = f"='Annual IS'!{ic}93"
        dc[f"{c}29"] = f"='Annual IS'!{ic}55"
        dc[f"{c}32"] = f"={c}6/{c}4"
        dc[f"{c}33"] = 0.065
        dc[f"{c}35"] = f"={c}11/{c}4"
        dc[f"{c}36"] = f"=IF({c}13>0,{c}15/{c}13,0)"
        dc[f"{c}38"] = None
        for ref in ["6", "9", "10", "11", "15", "18", "29"]:
            setfont(dc, f"{c}{ref}", GREEN if ref in ("6", "15", "18", "29") else BLACK)
        setfont(dc, f"{c}33", BLUE)
        if p: dc[f"{c}34"] = f"={c}10/{p}10-1"
        else: dc[f"{c}34"] = None
    else:
        dc[f"{c}32"] = f"=1-{RBT(c, 'gm')}"; setfont(dc, f"{c}32", GREEN)
        dc[f"{c}9"] = f"={RBT(c, 'svc')}*{c}33"
        dc[f"{c}10"] = f"={p}10*(1+{c}34)"
        dc[f"{c}11"] = f"={c}4*{c}35"
        dc[f"{c}15"] = f"=MAX({c}13,0)*{c}36"
        dc[f"{c}18"] = f"={c}11*{c}38"
        for ref in ["9", "10", "11", "15", "18"]:
            setfont(dc, f"{c}{ref}", BLACK)
    for rr in ["32", "33", "34", "35", "38"]:
        dc[f"{c}{rr}"].number_format = "0" if rr == "38" else PCT

# historical D&A, NWC, capex from the FY25 10-K cash flow statement ($mm)
hist_cf = {"C": (8.020, -0.095, 4.802), "D": (9.126, -15.220, 4.451), "E": (8.529, 3.355, 5.914)}
for c, (da, nwc, capex) in hist_cf.items():
    dc[f"{c}17"] = da; dc[f"{c}19"] = nwc; dc[f"{c}20"] = -capex
    for ref in ["17", "19", "20"]: setfont(dc, f"{c}{ref}", BLUE)
note(dc["C17"], "FY25 10-K cash flow statement: D&A $8.0M / $9.1M / $8.5M (FY23-25).")
note(dc["C19"], "FY25 10-K: sum of changes in operating assets & liabilities, excluding operating-lease liabilities (offset by non-cash lease expense).")
note(dc["C20"], "FY25 10-K: purchases of PP&E + capitalized internal-use software.")
note(dc["E29"], "FY23-25 weighted diluted shares are pre-IPO (IPO Sep 2025); forecasts grow from current shares outstanding (Key Stats).")

dc["F29"] = "='WACC Calculations'!C27*(1+F41)"; setfont(dc, "F29", GREEN)

fc_in = {33: [0.065] * 5, 34: [0.08, 0.04, 0.04, 0.04, 0.04], 35: [0.12, 0.10, 0.08, 0.07, 0.06],
         36: [0.21] * 5, 37: [0.018] * 5, 38: [0] * 5, 39: [-0.02] * 5, 40: [0.014] * 5, 41: [0.025] * 5}
for rr, vals in fc_in.items():
    for c, v in zip(FCST, vals):
        dc[f"{c}{rr}"] = v
note(dc["F33"], "Thesis 2.1: insurance ~3.5% (Bleecker: 3-4 pts of GM) + customer support ~3% of services revenue, both booked in G&A (10-K). Scales with hours.")
note(dc["F34"], "Other G&A + S&M + R&D ex-SBC. FY26 +8%: 1H26 annualised opex ex-SBC ($239M incl. hour-driven G&A) is ~10% above FY25 (Quarterly IS); fixed part ~8%. FY27+: 4% (consensus total opex +3.9%).")
note(dc["F35"], "1H26 SBC = 12% of revenue (Cash Flow tab), inflated by IPO grants; falls to 6% by FY30 as they vest.")
note(dc["F36"], "US federal rate on positive EBIT only; NOLs ignored (conservative for a short).")
note(dc["F38"], "0 = stock comp treated as a real cost (not added back). Set to 1 to add it back as the template originally did.")
note(dc["F39"], "Receivables grow with revenue (AR +$23.7M in 1H26).")
note(dc["F40"], "FY23-25 capex 1.4-1.9% of revenue (10-K).")
note(dc["F41"], "SBC dilution (~$60M/yr SBC on a ~$2.3B market cap).")

# memo: adj. EBITDA vs consensus
memo = {43: "Memo: Adj. EBITDA (EBIT + D&A + SBC)", 44: "   Consensus EBITDA (BBG; FY28 CapIQ)", 45: "   Model vs consensus ($mm)"}
for rr, t in memo.items():
    dc[f"B{rr}"] = t
    dc[f"B{rr}"]._style = copy(dc["B33"]._style)
for c in ALL:
    dc[f"{c}43"] = f"={c}13+{c}17+{c}11"
    dc[f"{c}43"]._style = copy(dc["C13"]._style); dc[f"{c}43"].number_format = "#,##0_);(#,##0)"
    if c in KS_COL:
        dc[f"{c}44"] = f"='Consensus (BBG)'!{BBG_COL[c]}10" if c in BBG_COL else f"='Key Stats'!{KS_COL[c]}22"
        dc[f"{c}45"] = f"={c}43-{c}44"
        for rr in (44, 45):
            dc[f"{c}{rr}"]._style = copy(dc["C12"]._style); dc[f"{c}{rr}"].number_format = "#,##0_);(#,##0)"
        setfont(dc, f"{c}44", GREEN)

# valuation panel: Via balance sheet and share count
dc["M14"] = "='Key Stats'!B63"; setfont(dc, "M14", GREEN)
dc["M15"] = 0; setfont(dc, "M15", BLUE)
note(dc["M15"], "No borrowings (10-Q Q2'26: $100M revolver undrawn apart from $28M letters of credit). CapIQ 'Total Debt' of $18.0M is operating lease liabilities ($9.8M current + $8.2M noncurrent); rent is already in opex, so it is not deducted again.")
dc["M18"] = "='WACC Calculations'!C27"; setfont(dc, "M18", GREEN)
dc["M31"] = "=SUM(F28:J28)"   # was SUM(F27:J27): undiscounted cash flows

# terminal value: normalized UFCF margin (Gordon) and EV/Revenue exit multiple, because
# 2030 UFCF and EBIT are still negative and cannot be capitalized
dc["L3"] = "Terminal-Year UFCF (normalized)"
dc["M3"] = "=J4*M6"
dc["L6"] = "Terminal UFCF Margin (normalized)"
dc["L6"]._style = copy(dc["L4"]._style)
dc["M6"] = 0.10
dc["M6"]._style = copy(dc["M4"]._style)
note(dc["M6"], "Steady-state UFCF margin after SBC: ~16% adj. EBITDA (below mgmt's 20-25% target) less ~4% SBC and ~1.4% capex, plus some working capital; taxes shielded by NOLs. 2030 UFCF is still negative, so it cannot be capitalized directly.")
dc["L26"] = "Last Year Revenue"
dc["M26"] = "=J4"
dc["L27"] = "Terminal EV/Revenue Multiple"
dc["M27"] = 2.5
dc["M27"].number_format = MULT
note(dc["M27"], "De-rating from 2.9x FY27E EV/revenue today (Key Stats) toward a services mix: ~20% software at ~6x + ~80% services at ~1.5x = ~2.5x. 2030 EBIT is negative, so an EBIT multiple cannot be used.")

# WACC tab: Via price and shares
wacc["C15"] = "='Key Stats'!B50"; setfont(wacc, "C15", GREEN)
wacc["C16"] = "='Key Stats'!B51"; setfont(wacc, "C16", GREEN)
# diluted shares (treasury stock method), Q2'26 10-Q
for rr, lab in [(20, "Diluted Shares (treasury stock method)"), (21, "Basic shares outstanding (mm)"),
                (22, "Options outstanding (mm)"), (23, "Weighted-avg. exercise price ($)"),
                (24, "Net shares from options (mm)"), (25, "Unvested RSUs (mm)"),
                (26, "PSUs (excluded: vest at targets vs $46 IPO price)"), (27, "Diluted Shares Outstanding (mm)")]:
    wacc[f"B{rr}"] = lab
    wacc[f"B{rr}"]._style = copy(wacc[f"B{16 if rr not in (20, 27) else 13}"]._style)
    wacc[f"C{rr}"]._style = copy(wacc["C16"]._style)
wacc["B20"]._style = copy(wacc["B2"]._style)
wacc["C20"]._style = copy(wacc["C2"]._style)
wacc["C21"] = "=C16"
wacc["C22"] = 8.901695; setfont(wacc, "C22", BLUE)
wacc["C23"] = 13.94; setfont(wacc, "C23", BLUE); wacc["C23"].number_format = "$#,##0.00"
wacc["C24"] = "=C22*MAX(0,1-C23/C15)"
wacc["C25"] = 3.750862; setfont(wacc, "C25", BLUE)
wacc["C26"] = 0; setfont(wacc, "C26", BLUE)
wacc["C27"] = "=C21+C24+C25+C26"
wacc["C27"]._style = copy(wacc["C13"]._style); wacc["C27"].number_format = "#,##0.0"
for rr in (21, 22, 24, 25, 26):
    wacc[f"C{rr}"].number_format = "#,##0.00"
note(wacc["C16"], "10-Q Q2'26 cover: 77,659,430 Class A + 3,846,183 Class B shares outstanding as of Jul 31, 2026 = 81.51M (via Key Stats).")
note(wacc["C22"], "10-Q Q2'26, stock option activity: 8,901,695 options outstanding at Jun 30, 2026, weighted-avg. exercise price $13.94.")
note(wacc["C25"], "10-Q Q2'26, RSU activity: 3,750,862 unvested RSUs at Jun 30, 2026.")
note(wacc["C26"], "2,486,727 PSUs (CEO 2,051,945 + CFO 434,782) vest only at stock-price targets relative to the $46 IPO price; excluded while the stock is well below.")
note(wacc["C11"], "Check: pull VIA's beta from Bloomberg (IPO Sep 2025, so history is short).")

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT, R)
